from pathlib import Path
from openpyxl import Workbook, load_workbook
from datetime import datetime

EXCEL_FILE = Path(__file__).with_name("bank_data.xlsx")
python -m pip install openpyxlfrom openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime

EXCEL_FILE = Path(__file__).with_name("bank_data.xlsx")


def setup_excel():
    """Create the Excel workbook and its sheets if they do not exist."""
    if EXCEL_FILE.exists():
        wb = load_workbook(EXCEL_FILE)
    else:
        wb = Workbook()
        accounts_ws = wb.active
        accounts_ws.title = "Accounts"
        accounts_ws.append(["Account Number", "Name", "Mobile", "Balance"])

        tx_ws = wb.create_sheet("Transactions")
        tx_ws.append(["Date & Time", "Account Number", "Transaction Type",
                      "Amount (Rs.)", "Balance After Transaction", "Status"])
        wb.save(EXCEL_FILE)

return wb, accounts_ws, tx_ws


# Set up Excel
wb, accounts_ws, tx_ws = setup_excel()


# Create a new account
def create_account():
    account_number = input("Enter Account Number: ").strip()

    # Check for duplicate account numbers
    for row in accounts_ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == account_number:
            print("Account number already exists!")
            return

    name = input("Enter Name: ").strip()
    mobile = input("Enter Mobile Number: ").strip()
    balance = 0

    if not account_number or not name or not mobile:
        print("All fields are required!")
        return

    accounts_ws.append([
        account_number, name, mobile, balance
    ])

    wb.save(EXCEL_FILE)
    print("Account created successfully!")


# Check account balance
def check_balance():
    account_number = input("Enter Account Number: ").strip()

    for row in accounts_ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == account_number:
            print("Account Holder:", row[1])
            print("Mobile Number:", row[2])
            print("Available Balance: Rs.", row[3] or 0)
            return

    print("Account not found!")


# Deposit money
def deposit():
    account_number = input("Enter Account Number: ").strip()

    for row in range(2, accounts_ws.max_row + 1):
        if str(accounts_ws.cell(row, 1).value) == account_number:
            try:
                amount = float(input("Enter Deposit Amount: Rs. "))

                if not 0 < amount < float("inf"):
                    print("Please enter a valid amount!")
                    return

            except ValueError:
                print("Invalid amount! Please enter a number.")
                return

            balance = float(accounts_ws.cell(row, 4).value or 0)
            balance = round(balance + amount, 2)

            accounts_ws.cell(row, 4).value = balance

            from datetime import datetime

            tx_ws.append([
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                account_number,
                "Deposit",
                amount,
                balance,
                "Success"
            ])

            wb.save(EXCEL_FILE)

            print("Money deposited successfully!")
            print("New Balance: Rs.", balance)
            return

    print("Account not found!")


# Main menu
while True:
    print("\n===== BANKING MANAGEMENT SYSTEM =====")
    print("1. Create Account")
    print("2. Check Balance")
    print("3. Deposit Money")
    print("4. Exit")

    choice = input("Enter your choice (1-4): ").strip()

    if choice == "1":
        create_account()

    elif choice == "2":
        check_balance()

    elif choice == "3":
        deposit()

    elif choice == "4":
        wb.save(EXCEL_FILE)
        print("All data saved successfully!")
from pathlib import Path
from openpyxl import Workbook, load_workbook
from datetime import datetime

EXCEL_FILE = Path(__file__).with_name("bank_data.xlsx")
        print("Thank you for using our Banking System!")
        break

    else:
        print("Invalid choice! Please try again.")
    if "Accounts" not in wb.sheetnames:
        ws = wb.create_sheet("Accounts", 0)
        ws.append(["Account Number", "Name", "Mobile", "Balance"])
    if "Transactions" not in wb.sheetnames:
        ws = wb.create_sheet("Transactions")
        ws.append(["Date & Time", "Account Number", "Transaction Type",
                   "Amount (Rs.)", "Balance After Transaction", "Status"])
    wb.save(EXCEL_FILE)
    return wb


def load_accounts():
    wb = setup_excel()
    ws = wb["Accounts"]
    accounts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        account_no, name, mobile, balance = row[:4]
        if account_no is not None:
            accounts[str(account_no)] = {
                "name": str(name or ""),
                "mobile": str(mobile or ""),
                "balance": float(balance or 0)
            }
    wb.close()
    return accounts


def save_account(account_no, item):
    wb = setup_excel()
    ws = wb["Accounts"]
    target_row = None
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 1).value) == str(account_no):
            target_row = row
            break
    values = [str(account_no), item["name"], item["mobile"], item["balance"]]
    if target_row is None:
        ws.append(values)
    else:
        for col, value in enumerate(values, start=1):
            ws.cell(target_row, col, value)
    wb.save(EXCEL_FILE)
    wb.close()


def log_transaction(account_no, transaction_type, amount, balance, status="Success"):
    wb = setup_excel()
    ws = wb["Transactions"]
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        str(account_no),
        transaction_type,
        float(amount),
        float(balance),
        status
    ])
    wb.save(EXCEL_FILE)
    wb.close()


def get_account_number():
    return input("Enter account number: ").strip()


def create_account(accounts):
    print("\n--- Create Account ---")
    account_no = get_account_number()
    if not account_no:
        print("Account number cannot be empty.")
        return
    if account_no in accounts:
        print("This account number already exists.")
        return

    name = input("Enter customer name: ").strip()
    mobile = input("Enter mobile number: ").strip()
    if not name or not mobile:
        print("Name and mobile number are required.")
        return

    accounts[account_no] = {"name": name, "mobile": mobile, "balance": 0.0}
    save_account(account_no, accounts[account_no])
    log_transaction(account_no, "Create Account", 0, 0)
    print("Account created successfully and saved to Excel!")


def deposit_money(accounts):
    print("\n--- Deposit Money ---")
    account_no = get_account_number()
    if account_no not in accounts:
        print("Account not found.")
        return
    try:
        amount = float(input("Enter deposit amount (Rs.): "))
        if amount <= 0:
            print("Amount must be greater than zero.")
            return
    except ValueError:
        print("Please enter a valid number.")
        return

    accounts[account_no]["balance"] += amount
    save_account(account_no, accounts[account_no])
    log_transaction(account_no, "Deposit", amount, accounts[account_no]["balance"])
    print(f"Deposit successful. Current balance: Rs. {accounts[account_no]['balance']:.2f}")


def withdraw_money(accounts):
    print("\n--- Withdraw Money ---")
    account_no = get_account_number()
    if account_no not in accounts:
        print("Account not found.")
        return
    try:
        amount = float(input("Enter withdrawal amount (Rs.): "))
        if amount <= 0:
            print("Amount must be greater than zero.")
            return
    except ValueError:
        print("Please enter a valid number.")
        return

    balance = accounts[account_no]["balance"]
    if amount > balance:
        print("Insufficient balance.")
        log_transaction(account_no, "Withdrawal", amount, balance, "Failed - Insufficient Balance")
        return

    accounts[account_no]["balance"] -= amount
    save_account(account_no, accounts[account_no])
    log_transaction(account_no, "Withdrawal", amount, accounts[account_no]["balance"])
    print(f"Withdrawal successful. Current balance: Rs. {accounts[account_no]['balance']:.2f}")


def show_balance(accounts):
    print("\n--- Show Balance ---")
    account_no = get_account_number()
    if account_no not in accounts:
        print("Account not found.")
        return
    item = accounts[account_no]
    print(f"Account Number: {account_no}")
    print(f"Name: {item['name']}")
    print(f"Mobile: {item['mobile']}")
    print(f"Current Balance: Rs. {item['balance']:.2f}")


def main():
    setup_excel()
    accounts = load_accounts()

    while True:
        print("\n====== BANK MANAGEMENT SYSTEM ======")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Show Balance")
        print("5. Exit")
        choice = input("Enter your choice (1-5): ").strip()

        if choice == "1":
            create_account(accounts)
        elif choice == "2":
            deposit_money(accounts)
        elif choice == "3":
            withdraw_money(accounts)
        elif choice == "4":
            show_balance(accounts)
        elif choice == "5":
            print("Thank you for using Bank Management System!")
            print(f"All account details and transaction history are saved in: {EXCEL_FILE}")
            break
        else:
            print("Invalid choice. Please enter a number from 1 to 5.")


if __name__ == "__main__":
    main()