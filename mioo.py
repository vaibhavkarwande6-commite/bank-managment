import os
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

FILE_NAME = "Bank_Data_Final.xlsx"

def get_wb():
    file = Path(FILE_NAME)
    if file.exists():
        return load_workbook(FILE_NAME)
    else:
        wb = Workbook()
        # Sheet 1 - Accounts
        ws1 = wb.active
        ws1.title = "Accounts"
        ws1.append(["Account Number", "Name", "Balance"])
        ws1.row_dimensions[1].font = Font(bold=True)

        # Sheet 2 - Transactions
        ws2 = wb.create_sheet("Transactions")
        ws2.append(["Date & Time", "Account Number", "Name", "Type", "Amount", "New Balance"])
        ws2.row_dimensions[1].font = Font(bold=True)
        wb.save(FILE_NAME)
        return wb

def save_transaction(acc_no, name, type_, amount, new_bal):
    wb = get_wb()
    ws = wb["Transactions"]
    now = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
    ws.append([now, acc_no, name, type_, amount, new_bal])
    wb.save(FILE_NAME)

def create_account():
    wb = get_wb()
    ws = wb["Accounts"]
    acc_no = input("Enter Account Number: ").strip()

    # check duplicate
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == acc_no:
            print("Account already exists!")
            return

    name = input("Enter Name: ").strip()
    balance = float(input("Enter Initial Balance: "))
    ws.append([acc_no, name, balance])
    wb.save(FILE_NAME)
    save_transaction(acc_no, name, "CREATE", balance, balance)
    print(f"Account Created & Saved in Excel Sheet!")

def deposit_money():
    wb = get_wb()
    ws = wb["Accounts"]
    acc_no = input("Enter Account Number: ").strip()
    found = False
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == acc_no:
            amount = float(input("Enter Amount to Deposit: "))
            row[2].value = float(row[2].value) + amount
            name = row[1].value
            new_bal = row[2].value
            wb.save(FILE_NAME)
            save_transaction(acc_no, name, "DEPOSIT", amount, new_bal)
            print(f"Deposited! New Balance: {new_bal}")
            found = True
            break
    if not found:
        print("Account not found!")

def withdraw_money():
    wb = get_wb()
    ws = wb["Accounts"]
    acc_no = input("Enter Account Number: ").strip()
    found = False
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == acc_no:
            amount = float(input("Enter Amount to Withdraw: "))
            if float(row[2].value) >= amount:
                row[2].value = float(row[2].value) - amount
                name = row[1].value
                new_bal = row[2].value
                wb.save(FILE_NAME)
                save_transaction(acc_no, name, "WITHDRAW", amount, new_bal)
                print(f"Withdrawn! New Balance: {new_bal}")
            else:
                print("Insufficient Balance!")
            found = True
            break
    if not found:
        print("Account not found!")

def show_balance():
    wb = get_wb()
    ws = wb["Accounts"]
    acc_no = input("Enter Account Number: ").strip()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == acc_no:
            print(f"Name: {row[1]}, Balance: {row[2]}")
            return
    print("Account not found!")

def main():
    while True:
        print("\n====== BANK MANAGEMENT SYSTEM ======")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Show Balance")
        print("5. Exit")
        choice = input("Enter your choice (1-5): ").strip()
        if choice == "1":
            create_account()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdraw_money()
        elif choice == "4":
            show_balance()
        elif choice == "5":
            print("All Data Saved in Bank_Data_Final.xlsx")
            break
        else:
            print("Invalid choice!")

if __name__ == "__main__":
    main()