from openpyxl import Workbook, load_workbook
from pathlib import Path

FILE = "bank_data.xlsx"
accounts = {}

def load_data():
    """Load accounts from Excel when the program starts."""
    if not Path(FILE).exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Accounts"
        ws.append(["Account Number", "Name", "Mobile", "Balance"])
        tx = wb.create_sheet("Transactions")
        tx.append(["Account Number", "Type", "Amount", "Balance After"])
        wb.save(FILE)
        return

    wb = load_workbook(FILE)
    ws = wb["Accounts"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        account_no, name, mobile, balance = row
        if account_no is not None:
            accounts[str(account_no)] = {
                "name": str(name),
                "mobile": str(mobile),
                "balance": float(balance)
            }
    wb.close()

def save_account(account_no):
    wb = load_workbook(FILE)
    ws = wb["Accounts"]
    found = False
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 1).value) == account_no:
            ws.cell(row, 2, accounts[account_no]["name"])
            ws.cell(row, 3, accounts[account_no]["mobile"])
            ws.cell(row, 4, accounts[account_no]["balance"])
            found = True
            break
    if not found:
        item = accounts[account_no]
        ws.append([account_no, item["name"], item["mobile"], item["balance"]])
    wb.save(FILE)
    wb.close()

def save_transaction(account_no, transaction_type, amount):
    wb = load_workbook(FILE)
    ws = wb["Transactions"]
    ws.append([account_no, transaction_type, amount, accounts[account_no]["balance"]])
    wb.save(FILE)
    wb.close()

def get_amount(prompt):
    try:
        amount = float(input(prompt))
        if amount <= 0:
            print("Amount must be greater than zero.")
            return None
        return amount
    except ValueError:
        print("Please enter a valid number.")
        return None

def create_account():
    account_no = input("Create account number: ").strip()
    if not account_no:
        print("Account number cannot be empty.")
        return
    if account_no in accounts:
        print("This account number already exists.")
        return

    name = input("Enter account holder name: ").strip()
    mobile = input("Enter mobile number: ").strip()
    if not name or not mobile:
        print("Name and mobile number are required.")
        return

    initial_balance = get_amount("Enter initial deposit (enter 1 or more): ")
    if initial_balance is None:
        return

    accounts[account_no] = {
        "name": name,
        "mobile": mobile,
        "balance": initial_balance
    }
    save_account(account_no)
    save_transaction(account_no, "Initial Deposit", initial_balance)
    print("Account created successfully!")
    print("Your account number is:", account_no)

def find_account():
    account_no = input("Enter account number: ").strip()
    if account_no not in accounts:
        print("Account not found. Please create an account first.")
        return None
    return account_no

def deposit_money():
    account_no = find_account()
    if account_no is None:
        return
    amount = get_amount("Enter deposit amount: Rs. ")
    if amount is None:
        return
    accounts[account_no]["balance"] += amount
    save_account(account_no)
    save_transaction(account_no, "Deposit", amount)
    print(f"Deposit successful. Balance: Rs. {accounts[account_no]['balance']:.2f}")

def withdraw_money():
    account_no = find_account()
    if account_no is None:
        return
    amount = get_amount("Enter withdrawal amount: Rs. ")
    if amount is None:
        return
    if amount > accounts[account_no]["balance"]:
        print("Insufficient balance.")
        return
    accounts[account_no]["balance"] -= amount
    save_account(account_no)
    save_transaction(account_no, "Withdrawal", amount)
    print(f"Withdrawal successful. Balance: Rs. {accounts[account_no]['balance']:.2f}")

def show_balance():
    account_no = find_account()
    if account_no is None:
        return
    item = accounts[account_no]
    print("\n--- Account Details ---")
    print("Account Number:", account_no)
    print("Name:", item["name"])
    print("Mobile:", item["mobile"])
    print(f"Current Balance: Rs. {item['balance']:.2f}")

def main():
    load_data()
    while True:
        print("\n====== BANK MANAGEMENT SYSTEM ======")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Show Balance")
        print("5. Exit")

        choice = input("Enter your choice (1-5): ").strip()
        if choice == "1":
            create_account()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdraw_money()
        elif choice == "4":
            show_balance()
        elif choice == "5":
            print("Thank you for using Bank Management System!")
            break
        else:
            print("Invalid choice. Please select 1 to 5.")

if __name__ == "__main__":
    main()